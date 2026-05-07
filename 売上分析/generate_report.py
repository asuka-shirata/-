import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from collections import defaultdict
import datetime

# ── データ読み込み ──────────────────────────────────────────
df = pd.read_csv("/Users/m0224/Desktop/ClaudeCode勉強会/売上分析/売上データ.csv")
df["日付"] = pd.to_datetime(df["日付"])
df["月"] = df["日付"].dt.to_period("M").astype(str)

# ── 集計 ────────────────────────────────────────────────────
cat_sales    = df.groupby("カテゴリ")["売上金額"].sum().sort_values(ascending=False)
month_sales  = df.groupby("月")["売上金額"].sum()
person_sales = df.groupby("担当者")["売上金額"].sum().sort_values(ascending=False)
cat_month    = df.pivot_table(index="月", columns="カテゴリ", values="売上金額", aggfunc="sum", fill_value=0)

# ── スタイル定数 ─────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="2F5496")
SUBHEADER_FILL= PatternFill("solid", fgColor="D6E4F7")
ALT_FILL      = PatternFill("solid", fgColor="F2F7FE")
TOTAL_FILL    = PatternFill("solid", fgColor="FFF2CC")
WHITE_FILL    = PatternFill("solid", fgColor="FFFFFF")
HEADER_FONT   = Font(name="メイリオ", bold=True, color="FFFFFF", size=11)
TITLE_FONT    = Font(name="メイリオ", bold=True, size=14, color="1F3864")
BODY_FONT     = Font(name="メイリオ", size=10)
BOLD_FONT     = Font(name="メイリオ", bold=True, size=10)
TOTAL_FONT    = Font(name="メイリオ", bold=True, size=10, color="7B3F00")

def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def thick_bottom():
    t = Side(style="medium", color="2F5496")
    s = Side(style="thin",   color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=t)

def set_header(ws, row, col, value, width=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font   = HEADER_FONT
    c.fill   = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thick_bottom()
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c

def set_cell(ws, row, col, value, fmt=None, bold=False, fill=None, align="right"):
    c = ws.cell(row=row, column=col, value=value)
    c.font   = BOLD_FONT if bold else BODY_FONT
    c.fill   = fill or WHITE_FILL
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.border = thin_border()
    if fmt:
        c.number_format = fmt
    return c

# ═══════════════════════════════════════════════════════════
wb = Workbook()

# ───────────────────────────────────────────────────────────
# シート1：サマリー
# ───────────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = "サマリー"
ws1.sheet_view.showGridLines = False
ws1.row_dimensions[1].height = 40
ws1.row_dimensions[2].height = 16

# タイトル
ws1.merge_cells("A1:G1")
t = ws1["A1"]
t.value     = "売上分析レポート　2024年1〜3月"
t.font      = Font(name="メイリオ", bold=True, size=16, color="1F3864")
t.fill      = PatternFill("solid", fgColor="EBF3FB")
t.alignment = Alignment(horizontal="center", vertical="center")

# 作成日
ws1["A2"].value = f"作成日：{datetime.date.today().strftime('%Y年%m月%d日')}"
ws1["A2"].font  = Font(name="メイリオ", size=9, color="808080")

# ── KPIボックス ───────────────────────────────────────────
kpi_data = [
    ("総売上", f"¥{df['売上金額'].sum():,}"),
    ("取引件数", f"{len(df):,} 件"),
    ("平均単価", f"¥{int(df['売上金額'].mean()):,}"),
    ("最高売上月", month_sales.idxmax()),
]
kpi_cols = [1, 3, 5, 7]
ws1.row_dimensions[4].height = 14
ws1.row_dimensions[5].height = 30
ws1.row_dimensions[6].height = 30
ws1.row_dimensions[7].height = 14

for i, (label, val) in enumerate(kpi_data):
    c = kpi_cols[i]
    ws1.merge_cells(start_row=5, start_column=c, end_row=5, end_column=c+1)
    ws1.merge_cells(start_row=6, start_column=c, end_row=6, end_column=c+1)
    lc = ws1.cell(row=5, column=c)
    lc.value = label
    lc.font = Font(name="メイリオ", size=9, color="2F5496", bold=True)
    lc.fill = PatternFill("solid", fgColor="D6E4F7")
    lc.alignment = Alignment(horizontal="center", vertical="center")
    lc.border = thin_border()
    vc = ws1.cell(row=6, column=c)
    vc.value = val
    vc.font = Font(name="メイリオ", size=14, bold=True, color="1F3864")
    vc.fill = PatternFill("solid", fgColor="EBF3FB")
    vc.alignment = Alignment(horizontal="center", vertical="center")
    vc.border = thin_border()

# ── カテゴリ別売上表 ──────────────────────────────────────
ws1.row_dimensions[9].height = 20
ws1.merge_cells("A9:C9")
h = ws1["A9"]
h.value = "カテゴリ別売上"
h.font  = TITLE_FONT
h.alignment = Alignment(horizontal="left", vertical="center")

headers = ["カテゴリ", "売上合計", "構成比"]
widths  = [20, 16, 12]
for i, (hd, wd) in enumerate(zip(headers, widths), 1):
    set_header(ws1, 10, i, hd)
    ws1.column_dimensions[get_column_letter(i)].width = wd

ws1.row_dimensions[10].height = 20
total = cat_sales.sum()
for ri, (cat, val) in enumerate(cat_sales.items(), 11):
    fill = ALT_FILL if ri % 2 == 0 else WHITE_FILL
    ws1.row_dimensions[ri].height = 18
    set_cell(ws1, ri, 1, cat,         align="left",  fill=fill)
    set_cell(ws1, ri, 2, val,         fmt="#,##0",   fill=fill)
    set_cell(ws1, ri, 3, val/total,   fmt="0.0%",    fill=fill)

tr = 11 + len(cat_sales)
ws1.row_dimensions[tr].height = 18
set_cell(ws1, tr, 1, "合計",    bold=True, fill=TOTAL_FILL, align="left")
set_cell(ws1, tr, 2, total,     fmt="#,##0", bold=True, fill=TOTAL_FILL)
set_cell(ws1, tr, 3, 1.0,       fmt="0.0%",  bold=True, fill=TOTAL_FILL)

# ── 月別売上表 ───────────────────────────────────────────
col_offset = 5
ws1.row_dimensions[9].height = 20
ws1.merge_cells(start_row=9, start_column=col_offset, end_row=9, end_column=col_offset+1)
mh = ws1.cell(row=9, column=col_offset, value="月別売上")
mh.font = TITLE_FONT
mh.alignment = Alignment(horizontal="left", vertical="center")

ws1.column_dimensions[get_column_letter(col_offset)].width = 12
ws1.column_dimensions[get_column_letter(col_offset+1)].width = 16
set_header(ws1, 10, col_offset,   "月")
set_header(ws1, 10, col_offset+1, "売上合計")

for ri, (mon, val) in enumerate(month_sales.items(), 11):
    fill = ALT_FILL if ri % 2 == 0 else WHITE_FILL
    ws1.row_dimensions[ri].height = 18
    set_cell(ws1, ri, col_offset,   mon, align="center", fill=fill)
    set_cell(ws1, ri, col_offset+1, val, fmt="#,##0",    fill=fill)

mtr = 11 + len(month_sales)
set_cell(ws1, mtr, col_offset,   "合計", bold=True, fill=TOTAL_FILL, align="center")
set_cell(ws1, mtr, col_offset+1, month_sales.sum(), fmt="#,##0", bold=True, fill=TOTAL_FILL)

# ───────────────────────────────────────────────────────────
# シート2：グラフ
# ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("グラフ")
ws2.sheet_view.showGridLines = False

ws2.merge_cells("A1:N1")
g = ws2["A1"]
g.value     = "売上分析グラフ"
g.font      = Font(name="メイリオ", bold=True, size=14, color="1F3864")
g.fill      = PatternFill("solid", fgColor="EBF3FB")
g.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 36

# ── グラフ用データ（隠し列として出力）───────────────────
data_start_row = 3

# カテゴリ
ws2.cell(row=data_start_row, column=17, value="カテゴリ")
ws2.cell(row=data_start_row, column=18, value="売上金額")
for i, (cat, val) in enumerate(cat_sales.items(), 1):
    ws2.cell(row=data_start_row+i, column=17, value=cat)
    ws2.cell(row=data_start_row+i, column=18, value=val)

# 月別
ws2.cell(row=data_start_row, column=20, value="月")
ws2.cell(row=data_start_row, column=21, value="売上金額")
for i, (mon, val) in enumerate(month_sales.items(), 1):
    ws2.cell(row=data_start_row+i, column=20, value=mon)
    ws2.cell(row=data_start_row+i, column=21, value=val)

# 担当者
ws2.cell(row=data_start_row, column=23, value="担当者")
ws2.cell(row=data_start_row, column=24, value="売上金額")
for i, (p, val) in enumerate(person_sales.items(), 1):
    ws2.cell(row=data_start_row+i, column=23, value=p)
    ws2.cell(row=data_start_row+i, column=24, value=val)

n_cat    = len(cat_sales)
n_month  = len(month_sales)
n_person = len(person_sales)

# ── グラフ1：カテゴリ別棒グラフ ─────────────────────────
bar = BarChart()
bar.type    = "col"
bar.title   = "カテゴリ別売上"
bar.y_axis.title = "売上金額（円）"
bar.style   = 10
bar.width   = 14
bar.height  = 10
bar.grouping = "clustered"

data_ref = Reference(ws2, min_col=18, min_row=data_start_row,
                     max_col=18, max_row=data_start_row+n_cat)
cats_ref = Reference(ws2, min_col=17, min_row=data_start_row+1,
                     max_row=data_start_row+n_cat)
bar.add_data(data_ref, titles_from_data=True)
bar.set_categories(cats_ref)
bar.series[0].graphicalProperties.solidFill = "2F5496"
ws2.add_chart(bar, "A3")

# ── グラフ2：月別折れ線グラフ ───────────────────────────
line = LineChart()
line.title   = "月別売上推移"
line.y_axis.title = "売上金額（円）"
line.style   = 10
line.width   = 14
line.height  = 10
line.smooth  = True

data_ref2 = Reference(ws2, min_col=21, min_row=data_start_row,
                      max_col=21, max_row=data_start_row+n_month)
cats_ref2 = Reference(ws2, min_col=20, min_row=data_start_row+1,
                      max_row=data_start_row+n_month)
line.add_data(data_ref2, titles_from_data=True)
line.set_categories(cats_ref2)
line.series[0].graphicalProperties.line.solidFill = "ED7D31"
line.series[0].graphicalProperties.line.width = 25000
ws2.add_chart(line, "H3")

# ── グラフ3：カテゴリ別円グラフ ─────────────────────────
pie = PieChart()
pie.title  = "カテゴリ別構成比"
pie.style  = 10
pie.width  = 14
pie.height = 10

data_ref3 = Reference(ws2, min_col=18, min_row=data_start_row,
                      max_col=18, max_row=data_start_row+n_cat)
cats_ref3 = Reference(ws2, min_col=17, min_row=data_start_row+1,
                      max_row=data_start_row+n_cat)
pie.add_data(data_ref3, titles_from_data=True)
pie.set_categories(cats_ref3)
pie.dataLabels = None
slice_colors = ["2F5496","ED7D31","A9D18E","FFC000"]
for i, color in enumerate(slice_colors[:n_cat]):
    pt = DataPoint(idx=i)
    pt.graphicalProperties.solidFill = color
    pie.series[0].dPt.append(pt)
ws2.add_chart(pie, "A22")

# ── グラフ4：担当者別棒グラフ ───────────────────────────
bar2 = BarChart()
bar2.type    = "bar"
bar2.title   = "担当者別売上"
bar2.x_axis.title = "売上金額（円）"
bar2.style   = 10
bar2.width   = 14
bar2.height  = 10

data_ref4 = Reference(ws2, min_col=24, min_row=data_start_row,
                      max_col=24, max_row=data_start_row+n_person)
cats_ref4 = Reference(ws2, min_col=23, min_row=data_start_row+1,
                      max_row=data_start_row+n_person)
bar2.add_data(data_ref4, titles_from_data=True)
bar2.set_categories(cats_ref4)
bar2.series[0].graphicalProperties.solidFill = "A9D18E"
ws2.add_chart(bar2, "H22")

# ── 隠し列を非表示 ──────────────────────────────────────
for c in range(16, 25):
    ws2.column_dimensions[get_column_letter(c)].hidden = True

# ───────────────────────────────────────────────────────────
# シート3：明細データ
# ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("明細データ")
ws3.sheet_view.showGridLines = False

ws3.merge_cells("A1:G1")
d = ws3["A1"]
d.value     = "売上明細データ"
d.font      = Font(name="メイリオ", bold=True, size=14, color="1F3864")
d.fill      = PatternFill("solid", fgColor="EBF3FB")
d.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 36

col_widths = [14, 22, 16, 10, 10, 14, 10]
for i, w in enumerate(col_widths, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

for i, col in enumerate(df.columns, 1):
    set_header(ws3, 2, i, col)
ws3.row_dimensions[2].height = 20

for ri, row in enumerate(df.itertuples(index=False), 3):
    fill = ALT_FILL if ri % 2 == 0 else WHITE_FILL
    ws3.row_dimensions[ri].height = 16
    set_cell(ws3, ri, 1, row.日付.strftime("%Y/%m/%d"), align="center", fill=fill)
    set_cell(ws3, ri, 2, row.商品名,    align="left",   fill=fill)
    set_cell(ws3, ri, 3, row.カテゴリ,  align="left",   fill=fill)
    set_cell(ws3, ri, 4, row.単価,      fmt="#,##0",     fill=fill)
    set_cell(ws3, ri, 5, row.数量,      fmt="#,##0",     fill=fill)
    set_cell(ws3, ri, 6, row.売上金額,  fmt="#,##0",     fill=fill)
    set_cell(ws3, ri, 7, row.担当者,    align="center",  fill=fill)

# 合計行
last = 3 + len(df)
ws3.row_dimensions[last].height = 18
set_cell(ws3, last, 1, "合計", bold=True, fill=TOTAL_FILL, align="center")
for c in [2,3,4,5]:
    set_cell(ws3, last, c, "", fill=TOTAL_FILL)
set_cell(ws3, last, 6, df["売上金額"].sum(), fmt="#,##0", bold=True, fill=TOTAL_FILL)
set_cell(ws3, last, 7, "", fill=TOTAL_FILL)

# フリーズ
ws3.freeze_panes = "A3"

# ── タブ色 ──────────────────────────────────────────────
ws1.sheet_properties.tabColor = "2F5496"
ws2.sheet_properties.tabColor = "ED7D31"
ws3.sheet_properties.tabColor = "A9D18E"

# 保存
out = "/Users/m0224/Desktop/ClaudeCode勉強会/売上分析/売上分析レポート.xlsx"
wb.save(out)
print(f"保存完了: {out}")
